"""
H46: lee un archivo con las columnas del catálogo de DIGEMID y valida cada fila.

No usa la BD: devuelve las filas válidas y los errores por fila, y importador.py
decide qué crear o actualizar. Los números de fila son los de Excel (encabezado = 1),
también para el CSV, para que el administrador encuentre la fila al abrir el archivo.

Si el archivo no se puede leer o le falta una columna obligatoria se lanza
ArchivoInvalido y no se importa nada.
"""
import csv
import io
import os
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException

EXTENSIONES = ('.xlsx', '.csv')

# campo -> (encabezado de DIGEMID, nombre con artículo para los mensajes, obligatorio, longitud máxima en la BD)
COLUMNAS = {
    'codigo': ('Código', 'el Código', False, 30),
    'nombre': ('Nombre Producto', 'el Nombre Producto', True, 150),
    'principio': ('Principio Activo', 'el Principio Activo', True, 150),
    'concentracion': ('Concentración', 'la Concentración', False, 50),
    'forma': ('Forma Farmacéutica', 'la Forma Farmacéutica', False, 100),
    'presentacion': ('Presentación', 'la Presentación', True, 100),
    'registro': ('Reg Sanitario', 'el Reg Sanitario', True, 50),
    'fecha': ('Fec Expiración', 'la Fec Expiración', False, None),
    'titular': ('Titular', 'el Titular', False, 200),
}
# Código y Titular pueden faltar en el archivo; las demás columnas tienen que estar
COLUMNAS_OPCIONALES = ('codigo', 'titular')


class ArchivoInvalido(Exception):
    """El archivo no se puede importar (formato, columnas que faltan): no se importa ninguna fila."""


@dataclass
class FilaDigemid:
    numero_fila: int
    codigo: str | None
    nombre: str
    principio: str
    concentracion: str | None
    forma: str | None
    presentacion: str
    registro: str
    fecha: date | None
    titular: str | None


@dataclass
class Lectura:
    filas: list[FilaDigemid] = field(default_factory=list)
    # (número de fila, registro sanitario o '', motivo)
    errores: list[tuple[int, str, str]] = field(default_factory=list)

    @property
    def filas_leidas(self):
        """Filas con datos (las totalmente vacías no cuentan)."""
        return len(self.filas) + len(self.errores)


def normalizar(texto):
    """Clave para comparar nombres sin distinguir mayúsculas, tildes ni espacios: LOSARTAN == Losartán."""
    sin_tildes = ''.join(c for c in unicodedata.normalize('NFKD', texto) if not unicodedata.combining(c))
    return ' '.join(sin_tildes.casefold().split())


def _clave_encabezado(texto):
    # "Fec. Expiración" = "FEC EXPIRACION"
    return normalizar(str(texto or '').replace('.', ' '))


ENCABEZADOS = {_clave_encabezado(encabezado): campo for campo, (encabezado, *_) in COLUMNAS.items()}


def leer(archivo, nombre):
    """Lee un .xlsx o .csv. `archivo` es una ruta o un archivo abierto en modo binario; `nombre` da la extensión."""
    extension = os.path.splitext(nombre)[1].lower()
    if extension not in EXTENSIONES:
        raise ArchivoInvalido(f'El archivo debe ser .xlsx o .csv (se recibió «{nombre}»).')
    if isinstance(archivo, (str, os.PathLike)):
        with open(archivo, 'rb') as abierto:
            contenido = abierto.read()
    else:
        contenido = archivo.read()
    if not contenido:
        raise ArchivoInvalido('El archivo está vacío.')

    filas = _filas_xlsx(contenido) if extension == '.xlsx' else _filas_csv(contenido)
    encabezado = next(filas, None)
    if encabezado is None:
        raise ArchivoInvalido('El archivo no tiene filas.')
    posiciones = _posiciones(encabezado)

    lectura = Lectura()
    primera_fila_por_registro = {}
    for numero_fila, valores in enumerate(filas, start=2):
        crudos = {campo: (valores[i] if i < len(valores) else None) for campo, i in posiciones.items()}
        if all(_texto(v) is None for v in crudos.values()):
            continue
        fila, motivos = _validar(numero_fila, crudos)
        if not motivos:
            primera = primera_fila_por_registro.get(fila.registro)
            if primera:
                motivos.append(f'registro sanitario {fila.registro} repetido (ya está en la fila {primera})')
            else:
                primera_fila_por_registro[fila.registro] = numero_fila
        if motivos:
            lectura.errores.append((numero_fila, fila.registro or '', '; '.join(motivos)))
        else:
            lectura.filas.append(fila)
    return lectura


def _filas_xlsx(contenido):
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise ArchivoInvalido('No se pudo abrir el Excel: el archivo no es un .xlsx válido.') from exc
    hoja = libro.worksheets[0]
    try:
        # Celdas y no solo valores: el Código necesita su formato para no perder ceros
        for celdas in hoja.iter_rows():
            yield [_valor_celda(c) for c in celdas]
    finally:
        libro.close()


def _valor_celda(celda):
    valor = getattr(celda, 'value', None)
    formato = getattr(celda, 'number_format', None) or ''
    # Un código guardado como número con formato "00000" se muestra con ceros en Excel: se respetan
    if isinstance(valor, int) and not isinstance(valor, bool) and formato and set(formato) == {'0'}:
        return str(valor).zfill(len(formato))
    return valor


def _filas_csv(contenido):
    try:
        texto = contenido.decode('utf-8-sig')
    except UnicodeDecodeError:
        # Excel en Windows guarda el CSV en cp1252 (Losartán con tilde)
        texto = contenido.decode('cp1252', errors='replace')
    primera_linea = texto.splitlines()[0] if texto else ''
    separador = ';' if primera_linea.count(';') > primera_linea.count(',') else ','
    yield from csv.reader(io.StringIO(texto, newline=''), delimiter=separador)


def _posiciones(encabezado):
    posiciones = {}
    for indice, valor in enumerate(encabezado):
        campo = ENCABEZADOS.get(_clave_encabezado(valor))
        if campo and campo not in posiciones:
            posiciones[campo] = indice
    faltan = [COLUMNAS[c][0] for c in COLUMNAS if c not in posiciones and c not in COLUMNAS_OPCIONALES]
    if faltan:
        raise ArchivoInvalido(
            'Faltan columnas obligatorias en la fila 1: ' + ', '.join(faltan) + '. '
            'No se importó nada.'
        )
    return posiciones


def _texto(valor):
    """Texto limpio (trim y espacios internos colapsados) o None si está vacío."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = ' '.join(str(valor).split())
    return texto or None


def _validar(numero_fila, crudos):
    motivos = []
    textos = {campo: _texto(valor) for campo, valor in crudos.items() if campo != 'fecha'}
    if textos.get('registro'):
        textos['registro'] = textos['registro'].upper()

    for campo, (encabezado, con_articulo, obligatorio, maximo) in COLUMNAS.items():
        if campo == 'fecha':
            continue
        valor = textos.get(campo)
        if valor is None:
            if obligatorio:
                genero = 'obligatoria' if con_articulo.startswith('la ') else 'obligatorio'
                motivos.append(f'falta {con_articulo} ({genero})')
        elif len(valor) > maximo:
            motivos.append(f'{encabezado} tiene {len(valor)} caracteres (máximo {maximo})')

    fecha, motivo_fecha = _fecha(crudos.get('fecha'))
    if motivo_fecha:
        motivos.append(motivo_fecha)

    fila = FilaDigemid(
        numero_fila=numero_fila,
        codigo=textos.get('codigo'),
        nombre=textos.get('nombre'),
        principio=textos.get('principio'),
        concentracion=textos.get('concentracion'),
        forma=textos.get('forma'),
        presentacion=textos.get('presentacion'),
        registro=textos.get('registro'),
        fecha=fecha,
        titular=textos.get('titular'),
    )
    return fila, motivos


def _fecha(valor):
    """(fecha, None) si es válida o vacía; (None, motivo) si no."""
    if isinstance(valor, datetime):
        return valor.date(), None
    if isinstance(valor, date):
        return valor, None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        # Número de serie de Excel en una celda sin formato de fecha
        try:
            return from_excel(valor).date(), None
        except (ValueError, OverflowError):
            return None, f'Fec Expiración «{valor}» no es una fecha válida'
    texto = _texto(valor)
    if texto is None:
        return None, None
    for formato in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(texto, formato).date(), None
        except ValueError:
            continue
    if _parece_fecha(texto):
        return None, f'fecha de expiración inválida ({texto} no existe)'
    return None, f'Fec Expiración «{texto}» no tiene el formato dd/mm/aaaa ni aaaa-mm-dd'


def _parece_fecha(texto):
    # 31/02/2029 tiene el formato correcto pero el día no existe
    partes = texto.replace('-', '/').split('/')
    return len(partes) == 3 and all(p.isdigit() for p in partes)
